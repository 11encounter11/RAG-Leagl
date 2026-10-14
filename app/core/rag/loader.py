"""
文档加载器 - 支持多种文档格式
使用 LlamaIndex 的 DocumentLoaders 加载文档
"""
import os
import logging
from pathlib import Path
from typing import List
from llama_index.core import Document
from llama_index.core.node_parser import SimpleNodeParser
from llama_index.readers.file import PyMuPDFReader

from app.config.settings import settings

logger = logging.getLogger(__name__)

"""
原始文件（如劳动法.pdf）
    ↓  Loader 加载
整篇纯文本（一个 Document 对象）
    ↓  Chunker 分块
多个小片段（多个 Node 对象）
    ↓  Embedder 向量化
多个向量
    ↓  存入 Milvus
向量数据库
"""
class DocumentLoader:
    """文档加载器 - 支持 PDF、DOCX、Excel、CSV、Markdown、TXT"""

    def __init__(self):
        self._readers = {
            ".pdf": PyMuPDFReader(),
        }

    def load_file(self, file_path: str) -> List[Document]:
        """
        加载单个文件
        Args:
            file_path: 文件路径
        Returns:
            Document 列表
        """
        ext = Path(file_path).suffix.lower()

        if ext == ".pdf":
            # PyMuPDFReader 的形参名是 file_path, 不是 file ——
            # 传 file= 会直接 TypeError(由 tests/test_loader.py::test_load_pdf 抓出)。
            documents = self._readers[".pdf"].load_data(file_path=Path(file_path))
        elif ext == ".docx":
            documents = self._load_docx(file_path)
        elif ext == ".xlsx":
            documents = self._load_excel(file_path)
        elif ext == ".csv":
            documents = self._load_csv(file_path)
        elif ext in (".md", ".txt"):
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            documents = [Document(text=content, metadata={"source": file_path})]
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

        # 为每个文档添加源文件元数据
        for doc in documents:
            doc.metadata["source_file"] = file_path
            doc.metadata["filename"] = Path(file_path).name

        return documents

    def _load_docx(self, file_path: str) -> List[Document]:
        """加载 Word 文档: 用 python-docx 直接读段落。

        刻意不用 llama_index 的 DocxReader —— 它依赖 docx2txt, 而 docx2txt
        并不在 requirements 里, 实际调用会在运行时抛 ImportError
        (由 tests/test_loader.py::test_load_docx 抓出)。
        这里改用已经在依赖清单里的 python-docx。
        """
        try:
            import docx

            document = docx.Document(file_path)
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
            if not text.strip():
                return []
            return [Document(text=text, metadata={"source": file_path})]
        except ImportError as e:  # pragma: no cover - 仅在依赖缺失时走到
            raise ValueError(f"Word 文件加载失败(缺少 python-docx): {e}") from e
        except Exception as e:
            raise ValueError(f"Word 文件加载失败: {e}") from e

    def _load_excel(self, file_path: str) -> List[Document]:
        """加载 Excel 文件，将每行转换为结构化文本"""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            documents = []
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                
                # 第一行作为表头
                headers = [str(cell) if cell else "" for cell in rows[0]]
                
                # 将每行数据转换为文本
                for row_idx, row in enumerate(rows[1:], start=2):
                    if not any(row):  # 跳过空行
                        continue
                    row_text = " | ".join(
                        f"{header}: {cell}" 
                        for header, cell in zip(headers, row) 
                        if cell is not None
                    )
                    if row_text.strip():
                        documents.append(Document(
                            text=row_text,
                            metadata={
                                "sheet": sheet_name,
                                "row": row_idx,
                                "source": file_path,
                            }
                        ))
            
            wb.close()
            return documents
        except Exception as e:
            raise ValueError(f"Excel 文件加载失败: {e}")

    def _load_csv(self, file_path: str) -> List[Document]:
        """加载 CSV 文件，将每行转换为结构化文本"""
        import csv
        documents = []
        
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row_idx, row in enumerate(reader, start=2):
                    row_text = " | ".join(
                        f"{key}: {value}" 
                        for key, value in row.items() 
                        if value
                    )
                    if row_text.strip():
                        documents.append(Document(
                            text=row_text,
                            metadata={
                                "row": row_idx,
                                "source": file_path,
                            }
                        ))
            return documents
        except Exception as e:
            raise ValueError(f"CSV 文件加载失败: {e}")

    def load_directory(self, dir_path: str) -> List[Document]:
        """
        加载目录下所有支持的文档
        Args:
            dir_path: 目录路径
        Returns:
            Document 列表
        """
        all_documents = []
        supported_ext = {".pdf", ".docx", ".xlsx", ".csv", ".md", ".txt"}

        for file_name in os.listdir(dir_path):
            file_path = os.path.join(dir_path, file_name)
            if Path(file_path).suffix.lower() in supported_ext:
                try:
                    docs = self.load_file(file_path)
                    all_documents.extend(docs)
                    logger.info("[DocumentLoader] 已加载: %s", file_name)
                except Exception as e:
                    logger.warning("[DocumentLoader] 加载失败 %s: %s", file_name, e)

        logger.info("[DocumentLoader] 共加载 %d 个文档", len(all_documents))
        return all_documents


class DocumentChunker:
    """文档分块器 - 将文档切分为适合检索的片段"""

    def __init__(
        self,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ):
        self.chunk_size = chunk_size or settings.CHUNK_SIZE
        self.chunk_overlap = chunk_overlap or settings.CHUNK_OVERLAP

        # 使用 LlamaIndex 的 SimpleNodeParser 进行分块
        self._parser = SimpleNodeParser.from_defaults(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
        )

    def chunk_documents(self, documents: List[Document]) -> list:
        """
        将文档列表切分为节点（chunks）
        Args:
            documents: Document 列表
        Returns:
            切分后的节点列表
        """
        nodes = self._parser.get_nodes_from_documents(documents)
        logger.info("[DocumentChunker] %d 个文档 → %d 个片段", len(documents), len(nodes))
        return nodes
